#!/usr/bin/env python3
"""Livestream comments -> order sheet for Trukuky.

Usage: python3 live_orders.py input.txt output.xlsx

Input sections (a header line in square brackets starts each section):
  [SẢN PHẨM]       code | name ("alias", "alias") | price | size:qty, size:qty
  [CÀI ĐẶT]        key = value  (ten_trang, ship, freeship_tu, thanh_toan,
                   tong_binh_luan, gio_het_live, qua_nua_dem)
  [MỚI NHẤT]       raw copy of the comments under the "Newest" filter
  [TẤT CẢ]         raw copy of the comments under the "All comments" filter
  [GIỜ CHÍNH XÁC]  C012 | 20:15[:30]       (optional, from hovering timestamps)
  [ĐẾM LẠI]        A1 | 90 | 1             (optional, physical recount)
  [SỬA]            C045 | A3 80 x1 ; A5 100  /  C046 | bỏ qua   (optional)

Design notes:
- The "Newest" view gives the order (reverse chronological); the "All comments"
  view gives completeness. Comments only present in "All" were filtered out of
  "Newest" by Facebook, so their position is unknown (ids start with L).
- Items are only auto-decided when unambiguous. Everything uncertain is
  flagged for a human instead of guessed.
"""
import math
import re
import sys
import unicodedata
from collections import Counter, OrderedDict, defaultdict
from dataclasses import dataclass, field
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from datetime import time as dtime


# ---------------------------------------------------------------- text utils

def strip_accents(s: str) -> str:
    s = s.replace('đ', 'd').replace('Đ', 'D')
    s = unicodedata.normalize('NFD', s)
    return ''.join(ch for ch in s if unicodedata.category(ch) != 'Mn')


def fold(s: str) -> str:
    """Lowercase, accent-free, single-spaced, for matching."""
    return re.sub(r'\s+', ' ', strip_accents(s).lower()).strip()


def clean_line(s: str) -> str:
    for ch in ('​', '‌', '‍', '‎', '‏', '﻿'):
        s = s.replace(ch, '')
    s = s.replace('\xa0', ' ')
    return re.sub(r'\s+', ' ', s).strip()


def vnd(n: Optional[int]) -> str:
    return '…đ' if n is None else f'{n:,.0f}'.replace(',', '.') + 'đ'


def parse_money(s: str) -> Optional[int]:
    t = fold(s).replace(' ', '')
    if not t:
        return None
    m = re.match(r'^(\d+(?:[.,]\d+)?)(k|nghin|ngan|tr|trieu|d|vnd|dong)?$', t)
    if not m:
        digits = re.sub(r'[^\d]', '', t)
        return int(digits) if digits else None
    num, unit = m.group(1), m.group(2)
    if unit in ('k', 'nghin', 'ngan'):
        return int(round(float(num.replace(',', '.')) * 1000))
    if unit in ('tr', 'trieu'):
        return int(round(float(num.replace(',', '.')) * 1_000_000))
    return int(num.replace('.', '').replace(',', ''))


# ------------------------------------------------------------ input sections

def split_sections(text: str) -> dict:
    sections, current = {}, None
    for raw in text.splitlines():
        m = re.match(r'^\s*\[(.+?)\]\s*$', raw)
        if m:
            current = fold(m.group(1))
            sections[current] = []
            continue
        if current is not None:
            sections[current].append(raw)
    return sections


def section(sections: dict, *names) -> list:
    for n in names:
        if fold(n) in sections:
            return sections[fold(n)]
    return []


# ------------------------------------------------------------------ products

SIZE_LETTERS = {'xs', 's', 'm', 'l', 'xl', 'xxl', '2xl', '3xl', 'freesize', 'fs'}
STOPWORDS = {
    'size', 'sz', 'shop', 'chot', 'lay', 'dat', 'mua', 'giu', 'cho', 'em', 'e',
    'chi', 'c', 'minh', 'nha', 'nhe', 'nhé', 'voi', 'va', 'cua', 'nay', 'kia',
    'ne', 'oi', 'ak', 'a', 'ah', 'ma', 'so', 'mau', 'cai', 'chiec', 'con', 'be',
    'khong', 'ko', 'k', 'co', 'duoc', 'them', 'x', 'sl', 'nua', 'di', 'luon',
}


@dataclass
class Product:
    code: str
    name: str
    price: Optional[int]
    variants: 'OrderedDict[str, Optional[int]]'
    aliases: list
    position: int

    def label(self) -> str:
        return f'{self.code} – {self.name}' if not self.code.startswith('SP') else self.name


def parse_products(lines: list) -> list:
    products = []
    for raw in lines:
        line = clean_line(raw)
        if not line or '|' not in line:
            continue
        parts = [p.strip() for p in line.split('|')]
        while len(parts) < 4:
            parts.append('')
        code, name, price, variants = parts[0], parts[1], parts[2], parts[3]
        if fold(code) in ('ma', 'code') or fold(name).startswith('ten'):
            continue  # header row
        aliases = re.findall(r'["“”]([^"“”]+)["“”]', name)
        aliases += re.findall(r'\(([^)]*)\)', name)
        base = re.sub(r'["“”][^"“”]+["“”]', '', name)
        base = clean_line(re.sub(r'\([^)]*\)', '', base))
        aliases = [a for a in (clean_line(x) for x in aliases)
                   if a and 'khong co ma' not in fold(a)]
        vmap = OrderedDict()
        for chunk in re.split(r'[,;]', variants):
            chunk = chunk.strip()
            if not chunk:
                continue
            m = re.match(r'^(.*?)\s*[:=]\s*(\d+)$', chunk)
            if m:
                vmap[m.group(1).strip()] = int(m.group(2))
            else:
                vmap[chunk] = None  # stock unknown
        if not vmap:
            vmap['freesize'] = None
        code = code or f'SP{len(products) + 1}'
        products.append(Product(code=code, name=base, price=parse_money(price),
                                variants=vmap, aliases=aliases,
                                position=len(products)))
    return products


def parse_settings(lines: list) -> dict:
    out = {}
    for raw in lines:
        m = re.match(r'^\s*([^=:]+?)\s*[=:]\s*(.+?)\s*$', raw)
        if m:
            out[fold(m.group(1)).replace(' ', '_')] = m.group(2)
    return out


# ------------------------------------------------------- facebook copy parser

NOISE_EXACT = {
    'thích', 'phản hồi', 'trả lời', 'chia sẻ', 'xem bản dịch', 'ẩn', 'xem thêm',
    'xem thêm bình luận', 'xem các bình luận trước', 'xem bình luận trước',
    'phù hợp nhất', 'mới nhất', 'tất cả bình luận', 'đã ghim', 'theo dõi',
    'viết bình luận...', 'viết bình luận…', 'viết bình luận', 'bình luận',
    'ẩn phản hồi', 'fan cứng', 'fan cuồng', 'người đóng góp nổi bật',
    'quản trị viên', 'người kiểm duyệt', 'xem bản gốc', 'đã gửi', 'gửi',
    'like', 'reply', 'share', 'see translation', 'hide', 'see more',
    'view more comments', 'view previous comments', 'most relevant', 'newest',
    'all comments', 'pinned', 'follow', 'write a comment...', 'write a comment…',
    'write a comment', 'comments', 'hide replies', 'top fan', 'top contributor',
    'admin', 'moderator', 'see original', '·', '…', '...',
}
NOISE_RE = [
    re.compile(r'^xem (tất cả )?\d+ (câu )?(phản hồi|trả lời)'),
    re.compile(r'^\d+ (phản hồi|trả lời|bình luận)$'),
    re.compile(r'^view (all )?\d+ (more )?(repl|comment)'),
    re.compile(r'^\d+ (repl(y|ies)|comments?)$'),
    re.compile(r'^.+ đã (trả lời|phản hồi)( ·)?( \d+ (phản hồi|trả lời))?$'),
    re.compile(r'^.+ replied( ·)?( \d+ repl(y|ies))?$'),
    re.compile(r'^\d+ (lượt thích|cảm xúc|reactions?)$'),
]
REPLY_MARKER_RE = re.compile(r'^(xem (tất cả )?\d+ (câu )?(phản hồi|trả lời)|ẩn phản hồi|'
                             r'view (all )?\d+ (more )?repl|hide replies|.+ đã (trả lời|phản hồi))')
AUTHOR_BADGES = {'tác giả', 'author'}
EDITED_MARKS = {'đã chỉnh sửa', 'edited'}
HIDDEN_RE = re.compile(r'^((bình luận )?(này )?đã (bị )?ẩn|hidden|this comment is hidden)')
TIME_RE = re.compile(
    r'^(vừa xong|just now|'
    r'\d+\s*(giây|phút|giờ|ngày|tuần|tháng|năm|g|ph|p|ng|tu|th|s|m|h|d|w|y|'
    r'min|mins|hr|hrs|sec|secs)( trước)?|'
    r'\d{1,2}:\d{2}(:\d{2})?|'
    r'(hôm qua|hôm nay|yesterday|today)( lúc| at)? \d{1,2}:\d{2}( ?[ap]m)?|'
    r'\d{1,2} (tháng|thg) \d{1,2}( lúc \d{1,2}:\d{2})?)$'
)
BADGE_SUFFIX_RE = re.compile(r'\s*·\s*(fan cứng|top fan|tác giả|author|theo dõi|follow)\b.*$',
                             re.IGNORECASE)


@dataclass
class Comment:
    name: str
    message: str = ''
    time_text: str = ''
    edited: bool = False
    author: bool = False
    hidden: bool = False
    view: str = ''
    paste_pos: int = 0            # position inside its own paste
    cid: str = ''
    chrono: Optional[int] = None  # 0 = earliest; None = unknown (filtered)
    only_in: str = ''             # '', 'all', 'newest'
    exact_time: Optional[int] = None
    # extraction results
    intent: str = ''
    items: list = field(default_factory=list)
    phones: list = field(default_factory=list)
    notes: list = field(default_factory=list)
    candidates: list = field(default_factory=list)
    inferred: bool = False
    answered: bool = False
    acked: bool = False
    manual: bool = False
    ignored: bool = False
    reply_marker: bool = False


def parse_fb_copy(lines: list, view: str) -> list:
    comments, cur, state, marker = [], None, 'seek', False
    for raw in lines:
        line = clean_line(raw)
        if not line:
            continue
        low = line.lower()
        if TIME_RE.match(low):
            if cur is not None:
                cur.time_text = line
                comments.append(cur)
                cur, state = None, 'post'
            continue
        if low in EDITED_MARKS:
            target = cur if cur is not None else (comments[-1] if comments else None)
            if target:
                target.edited = True
            continue
        if HIDDEN_RE.match(low):
            target = cur if cur is not None else (comments[-1] if comments else None)
            if target:
                target.hidden = True
            continue
        if low in AUTHOR_BADGES:
            if cur is not None:
                cur.author = True
            continue
        if low in NOISE_EXACT or any(r.match(low) for r in NOISE_RE):
            if REPLY_MARKER_RE.match(low):
                marker = True
            continue
        if state == 'post' and cur is None and re.fullmatch(r'\d+', line):
            continue  # reaction count under the previous comment
        if cur is None:
            badge = BADGE_SUFFIX_RE.search(line)
            name = BADGE_SUFFIX_RE.sub('', line).strip()
            cur = Comment(name=name, view=view, reply_marker=marker)
            marker = False
            if badge and fold(badge.group(1)) in ('tac gia', 'author'):
                cur.author = True
            state = 'in'
        else:
            cur.message = f'{cur.message} / {line}' if cur.message else line
    if cur is not None and cur.message:
        comments.append(cur)  # last comment without a time line
    for i, c in enumerate(comments):
        c.paste_pos = i
    return comments


def merge_views(newest: list, allc: list, page_name: str) -> tuple:
    """Returns (merged list, stats). Order comes from Newest, completeness from All."""
    chrono = list(reversed(newest))
    for i, c in enumerate(chrono):
        c.chrono = i
        c.cid = f'C{i + 1:03d}'
    pool = defaultdict(list)
    for c in chrono:
        pool[(fold(c.name), fold(c.message))].append(c)
    only_all = []
    for c in allc:
        key = (fold(c.name), fold(c.message))
        if pool[key]:
            match = pool[key].pop(0)
            match.edited = match.edited or c.edited
            match.hidden = match.hidden or c.hidden
            match.author = match.author or c.author
        else:
            c.only_in = 'all'
            only_all.append(c)
    only_newest = [c for lst in pool.values() for c in lst] if allc else []
    for c in only_newest:
        c.only_in = 'newest'
    for i, c in enumerate(only_all):
        c.cid = f'L{i + 1:03d}'
    merged = chrono + only_all
    pn = fold(page_name) if page_name else ''
    for c in merged:
        if pn and fold(c.name) == pn:
            c.author = True
    stats = {
        'newest': len(newest), 'all': len(allc),
        'only_all': len(only_all), 'only_newest': len(only_newest) if allc else 0,
        'merged': len(merged),
    }
    return merged, stats


# ---------------------------------------------------------------- extraction

PHONE_RE = re.compile(r'(?<!\d)(?:\+?84|0)(?:[\s.\-]?\d){9}(?!\d)')
ORDER_RE = re.compile(r'\b(chot|lay|dat|mua|giu|order|gom|ship cho|them)\b')
CANCEL_RE = re.compile(r'\b(huy|khong lay|ko lay|k lay|kh lay|khong lay nua|bo don|thoi khong|thoi ko|thoi k)\b')
QUESTION_RE = re.compile(
    r'(\?|\b(gia|bao nhieu|bn|bnhieu|bnhiu|con (khong|ko|k|hang|size|mau|ko a|k a)|'
    r'size gi|mac size|mac vua|may kg|chat (gi|vai|lieu)|vai gi|co nong|ship (bao|may)|'
    r'phi ship|bao lau|co size|co mau|co ko|co khong|sao)\b)')
CONTACT_RE = re.compile(r'\b(ib|inbox|check ib|rep ib|nhan tin|xem tin)\b')
WEIGHT_RE = re.compile(r'(\d{1,2}(?:[.,]\d)?)\s?(kg|ky|ki|can)\b')
AGE_RE = re.compile(r'(\d{1,2})\s?(thang|tuoi)\b')
HEIGHT_RE = re.compile(r'(\d{2,3})\s?cm\b')
QTY_RE = re.compile(r'(?:\bx|\bsl|so luong)\s?([1-9]\d?)\b|(?<!\d)([1-9])\s?(bo|cai|chiec|set|dam|vay|ao|quan)\b')
REPLY_TEXT_RE = re.compile(r'^(da|ok|oke|okie|roi|dung|vang)\b|^[✅✔]|\b(c|chi|minh|ban|e|em|me|a|anh) oi\b')
BROADCAST_RE = re.compile(r'\b(ca nha|moi nguoi|cac me|cac chi|cac ban|mau tiep|tiep theo|len song|sale|flash)\b')
ACK_RE = re.compile(r'\b(da chot|chot roi|da giu|giu roi|da len don|ok|oke|okie|nhan|da nhan)\b|✅|✔')


def code_regex(code: str) -> re.Pattern:
    c = fold(code)
    m = re.match(r'^([a-z]+)\s*0*(\d+)$', c)
    if m:
        letters, num = m.group(1), m.group(2)
        return re.compile(rf'(?<![a-z0-9]){letters}[\s\-]?0*{num}(?![0-9])')
    if c.isdigit():
        return re.compile(rf'(?:\bma|\bso|#)\s?0*{int(c)}(?![0-9])')
    return re.compile(rf'(?<![a-z0-9]){re.escape(c)}(?![a-z0-9])')


class Matcher:
    def __init__(self, products: list):
        self.products = products
        self.code_res = {p.code: code_regex(p.code) for p in products if not p.code.startswith('SP')}
        self.tokens = {}
        df = Counter()
        for p in products:
            toks = set()
            for phrase in [p.name] + p.aliases:
                toks |= {t for t in re.findall(r'[a-z0-9]+', fold(phrase)) if t not in STOPWORDS}
            self.tokens[p.code] = toks
            df.update(toks)
        n = max(len(products), 1)
        self.weight = {t: math.log((n + 1) / df[t]) for t in df}

    def by_code(self, text_f: str) -> list:
        return [p for p in self.products if p.code in self.code_res and self.code_res[p.code].search(text_f)]

    def by_name(self, text_f: str) -> tuple:
        words = set(re.findall(r'[a-z0-9]+', text_f)) - STOPWORDS
        scored = []
        for p in self.products:
            score = sum(self.weight[t] for t in self.tokens[p.code] & words)
            if score > 0:
                scored.append((score, p))
        scored.sort(key=lambda x: (-x[0], x[1].position))
        if not scored:
            return None, []
        best = scored[0]
        second = scored[1][0] if len(scored) > 1 else 0
        if best[0] >= 0.7 and best[0] - second >= 0.4:
            return best[1], []
        return None, [p for s, p in scored[:3] if s >= best[0] - 0.4]

    @staticmethod
    def sizes_in(text_f: str, product: Product) -> list:
        tokens = set(re.findall(r'[a-z0-9]+', text_f))
        found = []
        for key in product.variants:
            kt = [t for t in re.findall(r'[a-z0-9]+', fold(key))]
            if not kt:
                continue
            ok = True
            for t in kt:
                if t.isdigit():
                    if not re.search(rf'(?<!\d){t}(?!\d)', text_f):
                        ok = False
                elif t in SIZE_LETTERS:
                    if t not in tokens and not re.search(rf'\b(size|sz|s)\s?{t}\b', text_f):
                        ok = False
                elif t not in tokens:
                    ok = False
            if ok:
                found.append(key)
        return found


def extract(c: Comment, m: Matcher):
    if c.author or c.manual or c.ignored:
        return
    raw = c.message
    t = fold(raw)
    c.phones = ['0' + re.sub(r'\D', '', p)[-9:] for p in PHONE_RE.findall(raw)]
    t_clean = PHONE_RE.sub(' ', t)
    for rx, label in ((WEIGHT_RE, 'kg'), (AGE_RE, None), (HEIGHT_RE, 'cm')):
        for g in rx.findall(t_clean):
            c.notes.append(f'{g[0]} {g[1]}' if label is None else f'{g[0]}{label}')

    if CANCEL_RE.search(t_clean):
        c.intent = 'HUY'
        return

    products = m.by_code(t_clean)
    by_name = None
    if not products:
        by_name, c.candidates = m.by_name(t_clean)
        if by_name:
            products = [by_name]
    qty_m = QTY_RE.search(t_clean)
    qty = int(qty_m.group(1) or qty_m.group(2)) if qty_m else 1
    qty = max(1, min(qty, 20))
    has_order_word = bool(ORDER_RE.search(t_clean))
    is_question = bool(QUESTION_RE.search(t_clean)) and not has_order_word

    if by_name and not has_order_word and not Matcher.sizes_in(t_clean, by_name):
        c.intent = 'KHAC'
        c.notes.append(f'quan tâm {by_name.code}')
        return
    if products and not is_question:
        c.intent = 'CHOT'
        for p in products:
            text_wo_code = m.code_res[p.code].sub(' ', t_clean) if p.code in m.code_res else t_clean
            sizes = Matcher.sizes_in(text_wo_code, p)
            if len(p.variants) == 1:
                sizes = list(p.variants)
            if len(sizes) == 1:
                c.items.append({'product': p, 'variant': sizes[0], 'qty': qty, 'status': ''})
            elif len(sizes) > 1:
                for s in sizes:
                    c.items.append({'product': p, 'variant': s, 'qty': 1, 'status': ''})
            else:
                c.items.append({'product': p, 'variant': None, 'qty': qty, 'status': 'THIEU_SIZE'})
        return
    if c.candidates and (has_order_word or not is_question):
        c.intent = 'CHOT'
        c.items.append({'product': None, 'variant': None, 'qty': qty, 'status': 'MO_HO',
                        'candidates': c.candidates})
        return
    if has_order_word:
        c.intent = 'CHOT'
        c.items.append({'product': None, 'variant': None, 'qty': qty, 'status': 'KHONG_RO_MAU',
                        'size_hint': re.findall(r'(?<!\d)(\d{2,3})(?!\d)', t_clean)})
        return
    if is_question or products:
        c.intent = 'HOI'
        c.items = [{'product': p} for p in products]
        return
    if c.phones or CONTACT_RE.search(t_clean):
        c.intent = 'LIEN_HE'
        return
    c.intent = 'KHAC'


def infer_from_context(merged: list, m: Matcher):
    """Comments like 'chốt size 90': guess the product on screen from neighbours."""
    strong = [(c.chrono, c.items[0]['product']) for c in merged
              if c.chrono is not None and c.intent == 'CHOT'
              and len({i['product'].code for i in c.items if i.get('product')}) == 1
              and not c.inferred]
    for c in merged:
        for it in c.items:
            if it.get('status') != 'KHONG_RO_MAU' or c.chrono is None:
                continue
            near = sorted((abs(pos - c.chrono), pos, p) for pos, p in strong if pos != c.chrono)[:6]
            near = [p for _, _, p in near if _ <= 15]
            if len(near) < 4:
                continue
            top, cnt = Counter(p.code for p in near).most_common(1)[0]
            if cnt >= 4:
                p = next(x for x in m.products if x.code == top)
                sizes = [s for s in p.variants if any(re.fullmatch(r'\d+', tok) and tok == h
                                                       for tok in re.findall(r'[a-z0-9]+', fold(s))
                                                       for h in it.get('size_hint', []))]
                if len(p.variants) == 1:
                    sizes = list(p.variants)
                it['product'] = p
                it['variant'] = sizes[0] if len(sizes) == 1 else None
                it['status'] = 'SUY_RA' if it['variant'] else 'THIEU_SIZE'
                it['inferred'] = True
                c.inferred = True


def mark_replies(newest_chrono: list):
    """Page replies appear right after the comment they answer in the Newest paste."""
    by_paste = sorted([c for c in newest_chrono], key=lambda x: x.paste_pos)
    for i, c in enumerate(by_paste):
        if c.author:
            continue
        j = i + 1
        while j < len(by_paste) and by_paste[j].author:
            reply = fold(by_paste[j].message)
            reply_like = REPLY_TEXT_RE.search(reply) or ACK_RE.search(reply)
            if by_paste[j].reply_marker or (reply_like and not BROADCAST_RE.search(reply)):
                c.answered = True
                if ACK_RE.search(reply):
                    c.acked = True
            j += 1


# -------------------------------------------------------------- manual input

def apply_exact_times(merged: list, lines: list, overnight: bool):
    by_id = {c.cid: c for c in merged}
    for raw in lines:
        m = re.match(r'^\s*([CL]\d+)\s*\|\s*(\d{1,2}):(\d{2})(?::(\d{2}))?', raw.strip(), re.I)
        if not m or m.group(1).upper() not in by_id:
            continue
        h, mi, s = int(m.group(2)), int(m.group(3)), int(m.group(4) or 0)
        if overnight and h < 12:
            h += 24
        by_id[m.group(1).upper()].exact_time = h * 3600 + mi * 60 + s


def apply_recount(products: list, lines: list) -> dict:
    counted = {}
    for raw in lines:
        parts = [p.strip() for p in raw.split('|')]
        if len(parts) < 3 or not parts[2].isdigit():
            continue
        for p in products:
            if fold(p.code) == fold(parts[0]):
                for key in p.variants:
                    if fold(key) == fold(parts[1]):
                        counted[(p.code, key)] = int(parts[2])
    return counted


def apply_corrections(merged: list, products: list, lines: list):
    by_id = {c.cid: c for c in merged}
    for raw in lines:
        parts = [p.strip() for p in raw.split('|', 1)]
        if len(parts) < 2 or parts[0].upper() not in by_id:
            continue
        c = by_id[parts[0].upper()]
        if fold(parts[1]) in ('bo qua', 'bo', 'skip'):
            c.ignored, c.intent = True, 'BO_QUA'
            continue
        c.manual, c.intent, c.items, c.inferred = True, 'CHOT', [], False
        for spec in parts[1].split(';'):
            toks = spec.split()
            if not toks:
                continue
            p = next((x for x in products if fold(x.code) == fold(toks[0])), None)
            if not p:
                continue
            qty = 1
            variant = None
            for tok in toks[1:]:
                if re.fullmatch(r'x\d+', tok, re.I):
                    qty = int(tok[1:])
                else:
                    variant = next((k for k in p.variants if fold(k) == fold(tok)), variant)
            if len(p.variants) == 1:
                variant = next(iter(p.variants))
            c.items.append({'product': p, 'variant': variant, 'qty': qty,
                            'status': '' if variant else 'THIEU_SIZE'})


# ---------------------------------------------------------------- allocation

@dataclass
class Demand:
    customer: str
    product: Product
    variant: str
    qty: int
    comments: list
    status: str = ''
    rank: int = 0
    requested: int = 0
    pending: bool = False
    flags: list = field(default_factory=list)


def build_demands(merged: list) -> list:
    """One demand per (customer, product, variant); repeated comments are merged."""
    demands = OrderedDict()
    ordered = sorted(merged, key=lambda c: (c.chrono is None, c.chrono if c.chrono is not None else c.paste_pos))
    for c in ordered:
        if c.author or c.ignored or c.intent != 'CHOT':
            continue
        cust = fold(c.name)
        for it in c.items:
            if not it.get('product') or not it.get('variant'):
                continue
            key = (cust, it['product'].code, it['variant'])
            if key in demands:
                d = demands[key]
                d.comments.append(c)
                if re.search(r'\bthem\b', fold(c.message)):
                    d.qty += it['qty']
                    d.flags.append(f'{c.cid}: lấy thêm')
                else:
                    d.flags.append(f'{c.cid}: bình luận lặp, đã gộp')
            else:
                demands[key] = Demand(cust, it['product'], it['variant'], it['qty'], [c])
    return list(demands.values())


def order_key(d: Demand, use_exact: bool):
    first = d.comments[0]
    chrono = first.chrono if first.chrono is not None else 10 ** 6
    if use_exact:
        return (first.exact_time, chrono, first.cid)
    return (chrono, first.cid)


def allocate(demands: list, products: list, recount: dict) -> tuple:
    groups = defaultdict(list)
    for d in demands:
        groups[(d.product.code, d.variant)].append(d)
    contested, stock_left = [], {}
    for p in products:
        for key, qty in p.variants.items():
            stock_left[(p.code, key)] = recount.get((p.code, key), qty)
    for gkey, ds in groups.items():
        stock = stock_left.get(gkey)
        use_exact = all(d.comments[0].exact_time is not None for d in ds)
        ds.sort(key=lambda d: order_key(d, use_exact))
        total = sum(d.qty for d in ds)
        remaining = stock
        for rank, d in enumerate(ds, 1):
            d.rank = rank
            first = d.comments[0]
            if first.chrono is None:
                d.flags.append('không có trong bản "Mới nhất"' +
                               (': đã xếp theo giờ chính xác' if use_exact else ': chưa rõ thứ tự'))
            if first.edited:
                d.flags.append('bình luận đã chỉnh sửa')
            if first.inferred:
                d.flags.append('mẫu suy ra từ bình luận xung quanh')
            if remaining is None:
                d.status = 'GIU'
                d.flags.append('chưa có số tồn')
            elif remaining >= d.qty:
                d.status = 'GIU'
                remaining -= d.qty
            elif remaining > 0:
                d.status = 'GIU_MOT_PHAN'
                d.flags.append(f'chỉ còn {remaining}/{d.qty}')
                d.requested, d.qty, remaining = d.qty, remaining, 0
            else:
                d.status = 'HET'
        if remaining is not None:
            stock_left[gkey] = remaining
        if stock is not None and total > stock:
            uncertain = any(d.comments[0].chrono is None or d.comments[0].edited
                            or d.comments[0].inferred for d in ds)
            contested.append({'key': gkey, 'stock': stock, 'demand': total,
                              'demands': ds, 'uncertain': uncertain, 'exact': use_exact})
            has_filtered = any(d.comments[0].chrono is None for d in ds)
            for d in ds:
                c0 = d.comments[0]
                if (has_filtered and not use_exact) or (c0.edited and d.status != 'HET'):
                    d.pending = True
        for d in ds:
            if d.status == 'HET' and any(c.acked for c in d.comments):
                d.flags.append('SHOP ĐÃ TRẢ LỜI "ĐÃ CHỐT" TRÊN LIVE: anh/chị quyết')
                d.pending = True
    return contested, stock_left


# ------------------------------------------------------------------ messages

def pay_text(settings: dict) -> str:
    p = fold(settings.get('thanh_toan', ''))
    if not p:
        return ''
    if 'cod' in p and ('ck' in p or 'chuyen khoan' in p):
        return ' (mình chọn thanh toán khi nhận hàng hoặc chuyển khoản nha)'
    if 'cod' in p:
        return ' (thanh toán khi nhận hàng)'
    if 'ck' in p or 'chuyen khoan' in p:
        return ' (chuyển khoản trước giúp shop nha)'
    return f' ({settings["thanh_toan"]})'


def ship_for(subtotal: int, settings: dict) -> Optional[int]:
    ship = parse_money(settings.get('ship', '')) if settings.get('ship') else None
    free_from = parse_money(settings.get('freeship_tu', '')) if settings.get('freeship_tu') else None
    if free_from is not None and subtotal >= free_from:
        return 0
    return ship


LETTER_ORDER = {'xs': 0, 's': 1, 'm': 2, 'l': 3, 'xl': 4, 'xxl': 5, '2xl': 5, '3xl': 6}


def size_rank(v: str) -> Optional[float]:
    f = fold(v).strip()
    if re.fullmatch(r'\d{2,3}', f):
        return float(f)
    return LETTER_ORDER.get(f)


def alternatives(product: Product, variant: str, stock_left: dict, products: list) -> list:
    """Same model one size up (never smaller: the child would not fit it), then other
    models in the same size, preferring the same garment type (first word of the name)."""
    want = size_rank(variant)
    alts = []
    for k in product.variants:
        have = size_rank(k)
        if k == variant or (stock_left.get((product.code, k)) or 0) <= 0 or want is None or have is None:
            continue
        step = 15 if want >= 10 else 1
        if want < have <= want + step:
            alts.append(f'{product.code} size {k} (rộng hơn một chút)')
    ptype = fold(product.name).split(' ')[0] if product.name else ''
    others = []
    for p in products:
        if p.code == product.code:
            continue
        for k in p.variants:
            if (stock_left.get((p.code, k)) or 0) > 0 and fold(k) == fold(variant):
                same_type = fold(p.name).split(' ')[0] == ptype
                others.append((not same_type, p.position, f'mẫu {p.label()} size {k}'))
                break
    alts += [t for _, _, t in sorted(others)]
    return alts[:3]


def quote(msg: str, n: int = 60) -> str:
    return msg if len(msg) <= n else msg[:n].rsplit(' ', 1)[0] + '…'


STATUS_TEXT = {'GIU': 'Giữ', 'GIU_MOT_PHAN': 'Giữ 1 phần', 'HET': 'Hết'}
PRIORITY = ['Chốt rõ', 'Có món hết hàng', 'Cần hỏi lại', 'Hết hàng', 'Chờ xác minh (tranh chấp)',
            'Khách có bình luận HỦY: kiểm tra', 'Hỏi chưa được trả lời', 'Để lại SĐT / nhờ inbox']


def build_customers(merged, demands, stock_left, products, settings) -> list:
    custs = OrderedDict()
    ordered = sorted(merged, key=lambda c: (c.chrono is None, c.chrono if c.chrono is not None else c.paste_pos))
    for c in ordered:
        if c.author or c.ignored or c.intent in ('KHAC', '', 'BO_QUA'):
            continue
        k = fold(c.name)
        cu = custs.setdefault(k, {'name': c.name, 'comments': [], 'phones': [], 'notes': [],
                                  'demands': [], 'clarify': [], 'questions': [], 'cancel': False})
        cu['comments'].append(c)
        cu['phones'] += [p for p in c.phones if p not in cu['phones']]
        cu['notes'] += [n for n in c.notes if n not in cu['notes'] and not n.startswith('quan tâm')]
        if c.intent == 'HUY':
            cu['cancel'] = True
        if c.intent == 'HOI' and not c.answered:
            cu['questions'].append(c)
        for it in c.items:
            if c.intent == 'CHOT' and (not it.get('product') or not it.get('variant')):
                cu['clarify'].append((c, it))
    for d in demands:
        custs[d.customer]['demands'].append(d)

    rows = []
    for k, cu in custs.items():
        held = [d for d in cu['demands'] if d.status in ('GIU', 'GIU_MOT_PHAN')]
        out = [d for d in cu['demands'] if d.status == 'HET']
        partial = [d for d in cu['demands'] if d.status == 'GIU_MOT_PHAN']
        tentative = [d for d in held if d.comments[0].inferred]
        pending = [d for d in cu['demands'] if d.pending]
        subtotal = sum((d.product.price or 0) * d.qty for d in held)
        ship = ship_for(subtotal, settings) if held else None
        total = subtotal + ship if (held and ship is not None) else None
        phone_txt = f' (shop đã có SĐT đuôi {cu["phones"][0][-3:]})' if cu['phones'] else ''
        ask_phone = '' if cu['phones'] else ', SĐT'

        lines = ['Dạ shop Trukuky chào mình ạ.']
        firm = [d for d in held if d not in tentative]
        if firm:
            lines.append('Shop xác nhận đơn mình chốt trên live:')
            for d in firm:
                lines.append(f'• {d.product.label()} – size {d.variant} – {d.qty} – {vnd((d.product.price or 0) * d.qty)}')
        for d in tentative:
            lines.append(f'Lúc mình bình luận "{quote(d.comments[0].message)}", shop đang giới thiệu mẫu '
                         f'{d.product.label()}, shop giữ tạm cho mình size {d.variant} ({vnd(d.product.price)}). '
                         f'Mình xác nhận giúp shop đúng mẫu này không ạ?')
        if held:
            ship_txt = 'freeship' if ship == 0 else f'ship {vnd(ship)}'
            lines.append(f'Tạm tính {vnd(subtotal)} + {ship_txt} = Tổng {vnd(total)}{pay_text(settings)}.')
        for d in partial:
            lines.append(f'Mẫu {d.product.label()} size {d.variant} mình đặt {d.requested}, shop chỉ còn {d.qty} '
                         f'nên đang giữ {d.qty} cho mình ạ 🙏')
        for d in out:
            alts = alternatives(d.product, d.variant, stock_left, products)
            alt_txt = f' Shop còn {", ".join(alts)}. Mình muốn shop giữ mẫu nào ạ?' if alts else ''
            lines.append(f'Shop rất tiếc {d.product.label()} size {d.variant} đã hết do có bạn chốt trước mình 🙏{alt_txt}')
        for c, it in cu['clarify']:
            if it.get('status') == 'MO_HO' and it.get('candidates'):
                opts = ' '.join(f'{"①②③"[i]} {p.label()} – {vnd(p.price)}' for i, p in enumerate(it['candidates'][:3]))
                lines.append(f'Mình có chốt "{quote(c.message)}", shop có vài mẫu gần giống: {opts}. Mình chọn mẫu nào ạ?')
            elif it.get('status') == 'THIEU_SIZE' and it.get('product'):
                p = it['product']
                avail = [s for s in p.variants if (stock_left.get((p.code, s)) or 0) > 0 or p.variants[s] is None]
                lines.append(f'Mẫu {p.label()} ({vnd(p.price)}) còn size {", ".join(avail) or "…"}. '
                             f'Bé mặc size nào ạ (hoặc cho shop xin cân nặng, chiều cao của bé)?')
            else:
                lines.append(f'Mình có bình luận "{quote(c.message)}", shop chưa rõ mình chốt mẫu nào. '
                             f'Mình cho shop xin mã hoặc mô tả mẫu (màu, kiểu) và size giúp shop nha?')
        for c in cu['questions']:
            ans = []
            for it in c.items:
                p = it.get('product')
                if p:
                    left = [s for s in p.variants if (stock_left.get((p.code, s)) or 0) > 0 or p.variants[s] is None]
                    ans.append(f'{p.label()} giá {vnd(p.price)}, còn size {", ".join(left) or "đã hết"}')
            answer = ('; '.join(ans) + '.') if ans else '[nhân viên điền câu trả lời].'
            lines.append(f'Trên live mình có hỏi "{quote(c.message)}", shop chưa kịp trả lời, shop xin lỗi mình nha. {answer}')
        if held and not cu['clarify'] and not tentative:
            if out or partial:
                lines.append(f'Mình xác nhận giúp shop và gửi tên người nhận{ask_phone} và địa chỉ để shop gửi '
                             f'các món đang giữ nha{phone_txt}.')
            else:
                lines.append(f'Mình gửi giúp shop tên người nhận{ask_phone} và địa chỉ để shop gửi hàng ngay nha{phone_txt}.')
                if cu['notes']:
                    lines.append(f'Shop sẽ kiểm tra lại size theo thông tin mình báo ({", ".join(cu["notes"])}) ạ.')
                else:
                    lines.append('Nếu mình chưa chắc size, cho shop xin cân nặng và chiều cao của bé để shop kiểm tra giúp ạ.')
        elif cu['clarify'] or tentative:
            lines.append('Mình trả lời giúp shop để shop giữ hàng và báo tổng tiền chính xác nha ạ.')
        elif out:
            lines.append('Shop xin lỗi mình vì sự bất tiện này ạ.')
        elif cu['questions']:
            lines.append('Mình cần shop tư vấn size thì cho shop xin cân nặng, chiều cao của bé ạ.')
        else:
            lines.append('Shop thấy mình có để lại thông tin trên live, mình cần shop tư vấn mẫu nào ạ?')

        if pending:
            status = 'Chờ xác minh (tranh chấp)'
        elif cu['cancel']:
            status = 'Khách có bình luận HỦY: kiểm tra'
        elif cu['clarify'] or tentative:
            status = 'Cần hỏi lại'
        elif out and not held:
            status = 'Hết hàng'
        elif out or partial:
            status = 'Có món hết hàng'
        elif held:
            status = 'Chốt rõ'
        elif cu['questions']:
            status = 'Hỏi chưa được trả lời'
        else:
            status = 'Để lại SĐT / nhờ inbox'

        flags = []
        if pending:
            flags.append('KHÔNG GỬI TIN trước khi xác minh xong ca tranh chấp')
        for d in cu['demands']:
            flags += [f'{d.product.code}-{d.variant}: {f}' for f in d.flags]
        if len(cu['phones']) > 1:
            flags.append('nhiều SĐT khác nhau: có thể 2 người trùng tên')
        items = []
        for d in cu['demands']:
            items.append({'code': d.product.code, 'name': d.product.name, 'variant': d.variant,
                          'qty': d.qty, 'price': d.product.price,
                          'status': STATUS_TEXT[d.status] + (' (suy ra)' if d.comments[0].inferred else '')})
        rows.append({
            'name': cu['name'], 'phones': ', '.join(cu['phones']),
            'items_txt': '; '.join(f'{i["code"]} {i["variant"]} x{i["qty"]} [{i["status"]}]' for i in items),
            'items': items, 'status': status, 'flags': ' | '.join(flags), 'notes': ', '.join(cu['notes']),
            'message': '\n'.join(lines),
            'first': min((c.chrono for c in cu['comments'] if c.chrono is not None), default=10 ** 6),
            'ids': ', '.join(c.cid for c in cu['comments']),
        })
    rows.sort(key=lambda r: (PRIORITY.index(r['status']) if r['status'] in PRIORITY else 99, r['first']))
    for i, r in enumerate(rows, 1):
        r['kid'] = f'K{i:03d}'
    return rows


# --------------------------------------------------------------------- excel

FONT = 'Arial'
FILL = {
    'Chốt rõ': 'E3F2E1', 'Có món hết hàng': 'FFF4D6', 'Cần hỏi lại': 'FFF4D6',
    'Hết hàng': 'FBE3E1', 'Chờ xác minh (tranh chấp)': 'F3E5F5',
    'Khách có bình luận HỦY: kiểm tra': 'FBE3E1',
    'Hỏi chưa được trả lời': 'E6EEF9', 'Để lại SĐT / nhờ inbox': 'E6EEF9',
}
INPUT_FILL = PatternFill('solid', fgColor='FFFF00')


def col(i: int) -> str:
    return get_column_letter(i)


def write_sheet(ws, headers, rows, widths, wrap_cols=(), fill_col=None, input_cols=()):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(name=FONT, bold=True)
        cell.fill = INPUT_FILL if cell.column in input_cols else PatternFill('solid', fgColor='DDDDDD')
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[col(i)].width = w
    for row in ws.iter_rows(min_row=2):
        color = FILL.get(row[fill_col - 1].value) if fill_col else None
        for cell in row:
            cell.font = Font(name=FONT)
            cell.alignment = Alignment(wrap_text=(cell.column in wrap_cols), vertical='top')
            if cell.column in input_cols:
                cell.fill = INPUT_FILL
            elif color:
                cell.fill = PatternFill('solid', fgColor=color)
    ws.freeze_panes = 'A2'


def to_time(text: str):
    m = re.match(r'^\s*(\d{1,2}):(\d{2})', text or '')
    return dtime(int(m.group(1)) % 24, int(m.group(2))) if m else None


def export(path, customers, contested, recount_list, merged, stats, settings, unanswered):
    wb = Workbook()

    # --- Hướng dẫn -------------------------------------------------------
    wg = wb.active
    wg.title = 'Hướng dẫn'
    guide = [
        ['HƯỚNG DẪN DÙNG BẢNG CHỐT ĐƠN LIVE'],
        [''],
        ['Ô NỀN VÀNG = ô nhân viên cần điền. Các ô khác tự tính hoặc do AI điền, không sửa.'],
        [''],
        ['1. Trang "Chốt đơn": gửi tin theo thứ tự từ trên xuống (nhóm "Chốt rõ" trước).'],
        ['   Copy cột "Tin nhắn soạn sẵn", gửi bằng nút "Nhắn tin" dưới bình luận của khách.'],
        ['   Facebook chỉ cho Trang gửi 1 tin cho mỗi bình luận, trong 7 ngày, cho đến khi khách trả lời:'],
        ['   đọc lại tin trước khi gửi vì gửi rồi không bổ sung được.'],
        ['2. Gửi xong ghi ngay "Giờ gửi tin". Khách trả lời thì ghi "Giờ khách trả lời" và chọn "Kết quả".'],
        ['   Cột "Phút chờ" tự tính từ giờ kết thúc live (trang "Cài đặt"), dùng để đo tốc độ chốt đơn.'],
        ['3. Khách có trạng thái "Chờ xác minh (tranh chấp)": CHƯA gửi tin.'],
        ['   Mở trang "Ca tranh chấp", làm theo cột "Việc cần làm" (rê chuột vào chữ giờ dưới bình luận'],
        ['   để thấy giờ chính xác), điền giờ vào cột vàng rồi gửi lại cho AI phân hàng lần cuối.'],
        ['4. Trang "Cần đếm lại": đếm hàng thật, điền số vào cột vàng, gửi lại cho AI.'],
        ['5. Sửa số lượng hoặc đơn giá ở trang "Chi tiết món" thì tổng tiền tự tính lại,'],
        ['   nhưng tin nhắn soạn sẵn KHÔNG tự đổi: nhớ sửa tay số tiền trong tin nhắn.'],
        [''],
        ['VÍ DỤ MỘT DÒNG ĐÃ ĐIỀN Ở TRANG "Chốt đơn":'],
        ['Giờ gửi tin', 'Phút chờ', 'Giờ khách trả lời', 'Kết quả', 'Nhân viên'],
        [dtime(21, 5), '(tự tính)', dtime(21, 20), 'Xác nhận', 'Lan'],
        [''],
        ['MÀU TRẠNG THÁI:'],
    ] + [[k] for k in PRIORITY]
    for r in guide:
        wg.append(r)
    for row in wg.iter_rows():
        for cell in row:
            cell.font = Font(name=FONT, bold=(cell.row in (1, 18, 22)))
            if isinstance(cell.value, dtime):
                cell.number_format = 'hh:mm'
            if cell.row >= 23 and cell.value in FILL:
                cell.fill = PatternFill('solid', fgColor=FILL[cell.value])
    for c in wg[19]:
        c.font = Font(name=FONT, bold=True)
    for c in wg[20]:
        c.fill = INPUT_FILL
    wg.column_dimensions['A'].width = 16
    for letter in 'BCDE':
        wg.column_dimensions[letter].width = 18

    # --- Cài đặt ----------------------------------------------------------
    wc = wb.create_sheet('Cài đặt')
    ship = parse_money(settings['ship']) if settings.get('ship') else None
    free_from = parse_money(settings['freeship_tu']) if settings.get('freeship_tu') else None
    end_time = to_time(settings.get('gio_het_live', ''))
    conf = [
        ['Thiết lập', 'Giá trị', 'Nguồn'],
        ['Phí ship (đ)', ship, 'Anh/chị cung cấp' if ship is not None else 'Chưa có: điền vào ô vàng'],
        ['Freeship từ (đ)', free_from, 'Anh/chị cung cấp' if free_from is not None else 'Để trống nếu không có freeship'],
        ['Giờ kết thúc live', end_time, 'Anh/chị cung cấp' if end_time else 'Chưa có: điền dạng 22:30 để tính "Phút chờ"'],
        ['Thanh toán', settings.get('thanh_toan', ''), 'Anh/chị cung cấp'],
    ]
    for r in conf:
        wc.append(r)
    for row in wc.iter_rows():
        for cell in row:
            cell.font = Font(name=FONT, bold=(cell.row == 1))
            if cell.row > 1 and cell.column == 2:
                cell.fill = INPUT_FILL
    wc['B2'].number_format = wc['B3'].number_format = '#,##0'
    wc['B4'].number_format = 'hh:mm'
    wc.column_dimensions['A'].width = 20
    wc.column_dimensions['B'].width = 14
    wc.column_dimensions['C'].width = 48

    # --- Chốt đơn ---------------------------------------------------------
    ws = wb.create_sheet('Chốt đơn', 1)
    n = len(customers)
    rows = []
    for i, r in enumerate(customers, 1):
        x = i + 1
        rows.append([
            i, r['kid'], r['name'], r['status'], r['items_txt'],
            f"=SUMIFS('Chi tiết món'!$H:$H,'Chi tiết món'!$A:$A,B{x})",
            f"=IF(F{x}=0,\"\",IF(AND('Cài đặt'!$B$3<>\"\",F{x}>='Cài đặt'!$B$3),0,"
            f"IF('Cài đặt'!$B$2=\"\",\"\",'Cài đặt'!$B$2)))",
            f"=IF(OR(F{x}=0,G{x}=\"\"),\"\",F{x}+G{x})",
            r['phones'], r['notes'], r['flags'], r['message'], r['ids'],
            None,
            f"=IF(AND(ISNUMBER(N{x}),ISNUMBER('Cài đặt'!$B$4)),ROUND(MOD(N{x}-'Cài đặt'!$B$4,1)*1440,0),\"\")",
            None, None, None,
        ])
    write_sheet(ws, ['STT', 'Mã khách', 'Tên Facebook', 'Trạng thái', 'Món (mã size xSL [kết quả])',
                     'Tạm tính (đ)', 'Ship (đ)', 'Tổng (đ)', 'SĐT trong bình luận', 'Thông tin bé',
                     'Cần kiểm tra', 'Tin nhắn soạn sẵn', 'Mã bình luận', 'Giờ gửi tin', 'Phút chờ (tự tính)',
                     'Giờ khách trả lời', 'Kết quả', 'Nhân viên'],
                rows, [5, 8, 20, 18, 28, 12, 9, 12, 14, 12, 34, 70, 12, 10, 10, 11, 14, 11],
                wrap_cols=(5, 11, 12), fill_col=4, input_cols=(14, 16, 17, 18))
    for x in range(2, n + 2):
        for letter in 'FGH':
            ws[f'{letter}{x}'].number_format = '#,##0'
        ws[f'N{x}'].number_format = ws[f'P{x}'].number_format = 'hh:mm'
        ws[f'O{x}'].number_format = '0'
    dv = DataValidation(type='list', formula1='"Xác nhận,Không trả lời,Hủy,Đổi mẫu"', allow_blank=True)
    ws.add_data_validation(dv)
    if n:
        dv.add(f'Q2:Q{n + 1}')

    # --- Chi tiết món -----------------------------------------------------
    wi = wb.create_sheet('Chi tiết món', 2)
    irows = []
    for r in customers:
        for it in r['items']:
            x = len(irows) + 2
            irows.append([r['kid'], r['name'], it['code'], it['name'], it['variant'], it['qty'],
                          it['price'], f'=IF(I{x}="Hết",0,F{x}*G{x})', it['status']])
    write_sheet(wi, ['Mã khách', 'Tên Facebook', 'Mã SP', 'Tên mẫu', 'Size', 'SL', 'Đơn giá (đ)',
                     'Thành tiền (đ)', 'Kết quả'], irows, [9, 20, 8, 28, 8, 6, 12, 13, 14],
                input_cols=(6, 7))
    for x in range(2, len(irows) + 2):
        wi[f'G{x}'].number_format = wi[f'H{x}'].number_format = '#,##0'

    # --- Ca tranh chấp ----------------------------------------------------
    ws2 = wb.create_sheet('Ca tranh chấp')
    rows2 = []
    for g in contested:
        code, variant = g['key']
        held_ranks = [x.rank for x in g['demands'] if x.status == 'GIU']
        last_held = max(held_ranks) if held_ranks else 0
        has_filtered = any(x.comments[0].chrono is None for x in g['demands'])
        for d in g['demands']:
            c = d.comments[0]
            todo = []
            if has_filtered and not g['exact']:
                todo.append('RÊ CHUỘT XEM GIỜ (nhóm có bình luận bị lọc)')
            elif d.rank in (last_held, last_held + 1) and not g['exact']:
                todo.append('nên xem giờ (ranh giới được/hết)')
            if c.edited and d.status != 'HET':
                todo.append('XEM LỊCH SỬ SỬA: bấm "Đã chỉnh sửa"')
            if c.inferred:
                todo.append('HỎI KHÁCH XÁC NHẬN MẪU')
            rows2.append([f'{code} size {variant}', g['stock'], g['demand'], d.rank, c.cid, c.name,
                          c.message, STATUS_TEXT[d.status], '; '.join(d.flags), '; '.join(todo),
                          dtime(c.exact_time // 3600 % 24, c.exact_time // 60 % 60) if c.exact_time is not None else None])
    write_sheet(ws2, ['Mẫu', 'Tồn', 'Số lượng chốt', 'Thứ tự', 'Mã BL', 'Tên Facebook', 'Bình luận',
                      'Kết quả tạm', 'Lý do cần xem', 'Việc cần làm', 'Giờ chính xác'],
                rows2, [13, 6, 9, 7, 7, 18, 34, 10, 34, 30, 12], wrap_cols=(7, 9, 10), input_cols=(11,))
    for x in range(2, len(rows2) + 2):
        ws2[f'K{x}'].number_format = 'hh:mm'

    # --- Cần đếm lại ------------------------------------------------------
    ws3 = wb.create_sheet('Cần đếm lại')
    write_sheet(ws3, ['Mã', 'Tên', 'Size', 'Tồn trên danh sách', 'Số lượng chốt', 'Lý do', 'Số đếm thực tế'],
                [r[:6] + [None] for r in recount_list], [8, 28, 8, 14, 12, 30, 14],
                wrap_cols=(2, 6), input_cols=(7,))

    ws4 = wb.create_sheet('Hỏi chưa trả lời')
    write_sheet(ws4, ['Mã BL', 'Tên Facebook', 'Câu hỏi'],
                [[c.cid, c.name, c.message] for c in unanswered], [8, 22, 70], wrap_cols=(3,))

    ws5 = wb.create_sheet('Bị lọc (chỉ ở Tất cả)')
    write_sheet(ws5, ['Mã BL', 'Tên Facebook', 'Bình luận', 'Phân loại'],
                [[c.cid, c.name, c.message, c.intent] for c in merged if c.only_in == 'all'],
                [8, 22, 60, 12], wrap_cols=(3,))

    ws6 = wb.create_sheet('Toàn bộ bình luận')
    rows6 = []
    for c in sorted(merged, key=lambda c: (c.chrono is None, c.chrono if c.chrono is not None else c.paste_pos)):
        items = '; '.join(
            (f'{i["product"].code} {i.get("variant") or "?"} x{i.get("qty", 1)}' if i.get('product') else
             'mơ hồ: ' + '/'.join(p.code for p in i.get('candidates', [])) if i.get('candidates') else '?')
            for i in c.items)
        flags = ', '.join(f for f, on in (('shop', c.author), ('đã sửa', c.edited), ('bị ẩn', c.hidden),
                                          ('bị lọc', c.only_in == 'all'), ('chỉ ở Mới nhất', c.only_in == 'newest'),
                                          ('suy ra', c.inferred), ('sửa tay', c.manual),
                                          ('shop đã trả lời', c.answered)) if on)
        rows6.append([c.cid, '' if c.chrono is None else c.chrono + 1, c.name, c.message, c.time_text,
                      c.intent or ('SHOP' if c.author else ''), items, ', '.join(c.phones), flags])
    write_sheet(ws6, ['Mã BL', 'Thứ tự', 'Tên Facebook', 'Bình luận', 'Giờ (Facebook)', 'Phân loại',
                      'Đọc được', 'SĐT', 'Ghi chú'], rows6, [8, 7, 20, 50, 10, 10, 22, 13, 30], wrap_cols=(4,))

    # --- Kiểm tra đầy đủ ---------------------------------------------------
    ws7 = wb.create_sheet('Kiểm tra đầy đủ')
    shown = settings.get('tong_binh_luan')
    shown_n = int(re.sub(r'\D', '', shown)) if shown and re.sub(r'\D', '', shown) else None
    last = n + 1
    rows7 = [
        ['Số bình luận đọc được ở "Mới nhất"', stats['newest'], 'Đếm từ bản copy'],
        ['Số bình luận đọc được ở "Tất cả bình luận"', stats['all'], 'Đếm từ bản copy'],
        ['Chỉ có ở "Tất cả" (Facebook lọc khỏi "Mới nhất")', stats['only_all'], 'Đếm từ bản copy'],
        ['Chỉ có ở "Mới nhất" (bản copy "Tất cả" có thể thiếu)', stats['only_newest'], 'Đếm từ bản copy'],
        ['Tổng sau khi gộp', stats['merged'], 'Đếm từ bản copy'],
        ['Tổng hiển thị dưới video', shown_n, 'Anh/chị ghi (ô vàng)'],
        ['Chênh lệch (dương = có thể còn sót)', '=IF(B7="","",B7-B6)', 'Tự tính'],
        ['Số khách trong bảng chốt đơn', f"=COUNTA('Chốt đơn'!B2:B{max(last, 2)})", 'Tự tính'],
        ['Số khách "Chốt rõ"', f"=COUNTIF('Chốt đơn'!D2:D{max(last, 2)},\"Chốt rõ\")", 'Tự tính'],
        ['Số ca tranh chấp', len(contested), 'Tính khi phân hàng'],
        ['Số khách đã gửi tin', f"=COUNT('Chốt đơn'!N2:N{max(last, 2)})", 'Tự tính'],
        ['Phút chờ trung vị (từ lúc hết live đến lúc gửi tin)',
         f"=IF(COUNT('Chốt đơn'!O2:O{max(last, 2)})=0,\"\",MEDIAN('Chốt đơn'!O2:O{max(last, 2)}))", 'Tự tính'],
        ['Tỷ lệ khách xác nhận / số đã gửi tin',
         f"=IF(B12=0,\"\",COUNTIF('Chốt đơn'!Q2:Q{max(last, 2)},\"Xác nhận\")/B12)", 'Tự tính'],
    ]
    if not stats['newest']:
        rows7.append(['CẢNH BÁO: thiếu bản copy "Mới nhất"', 'Thứ tự chưa biết',
                      'Mọi ca tranh chấp phải xem giờ chính xác'])
    if not stats['all']:
        rows7.append(['CẢNH BÁO: thiếu bản copy "Tất cả bình luận"', 'Chưa kiểm tra',
                      'Chưa phát hiện được bình luận bị lọc'])
    write_sheet(ws7, ['Chỉ số', 'Giá trị', 'Nguồn'], rows7, [52, 16, 40])
    ws7['B7'].fill = INPUT_FILL
    ws7['B14'].number_format = '0.0%'
    wb.move_sheet('Hướng dẫn', offset=-wb.index(wb['Hướng dẫn']))
    wb.save(path)


# ---------------------------------------------------------------------- main

def run(text: str, out_path: str) -> dict:
    sec = split_sections(text)
    products = parse_products(section(sec, 'SẢN PHẨM', 'SAN PHAM'))
    settings = parse_settings(section(sec, 'CÀI ĐẶT', 'CAI DAT', 'SHIP & THANH TOÁN'))
    newest = parse_fb_copy(section(sec, 'MỚI NHẤT', 'MOI NHAT'), 'newest')
    allc = parse_fb_copy(section(sec, 'TẤT CẢ', 'TAT CA', 'TẤT CẢ BÌNH LUẬN'), 'all')
    merged, stats = merge_views(newest, allc, settings.get('ten_trang', 'Trukuky'))
    matcher = Matcher(products)
    apply_corrections(merged, products, section(sec, 'SỬA', 'SUA'))
    for c in merged:
        extract(c, matcher)
    infer_from_context(merged, matcher)
    mark_replies([c for c in merged if c.chrono is not None])
    overnight = fold(settings.get('qua_nua_dem', '')) in ('co', 'yes', '1')
    apply_exact_times(merged, section(sec, 'GIỜ CHÍNH XÁC', 'GIO CHINH XAC'), overnight)
    recount = apply_recount(products, section(sec, 'ĐẾM LẠI', 'DEM LAI'))
    demands = build_demands(merged)
    contested, stock_left = allocate(demands, products, recount)

    recount_list = []
    for p in products:
        for key, qty in p.variants.items():
            dem = sum(d.qty for d in demands if d.product.code == p.code and d.variant == key)
            reasons = []
            if qty is not None and dem >= qty and dem > 0:
                reasons.append('số chốt ≥ tồn')
            if qty is not None and 0 < qty <= 2 and dem > 0:
                reasons.append('tồn ≤ 2')
            if qty is None and dem > 0:
                reasons.append('chưa có số tồn')
            if reasons and (p.code, key) not in recount:
                recount_list.append([p.code, p.name, key, qty, dem, ', '.join(reasons), ''])

    customers = build_customers(merged, demands, stock_left, products, settings)
    unanswered = [c for c in merged if c.intent == 'HOI' and not c.answered and not c.author]
    export(out_path, customers, contested, recount_list, merged, stats, settings, unanswered)
    return {'stats': stats, 'customers': customers, 'contested': contested,
            'recount': recount_list, 'unanswered': unanswered, 'merged': merged}


if __name__ == '__main__':
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    with open(sys.argv[1], encoding='utf-8') as fh:
        result = run(fh.read(), sys.argv[2])
    s = result['stats']
    print(f"Mới nhất: {s['newest']} | Tất cả: {s['all']} | bị lọc: {s['only_all']} | "
          f"chỉ ở Mới nhất: {s['only_newest']} | sau gộp: {s['merged']}")
    print(f"Khách: {len(result['customers'])} | ca tranh chấp: {len(result['contested'])} | "
          f"cần đếm lại: {len(result['recount'])} | hỏi chưa trả lời: {len(result['unanswered'])}")
    for r in result['customers']:
        print(f"- {r['kid']} {r['status']:<26} {r['name']:<12} {r['items_txt']}  {('⚠ ' + r['flags']) if r['flags'] else ''}")
